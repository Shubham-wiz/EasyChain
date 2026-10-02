"""Knowledge Bases: reading files, chunking, storage on SQLite and Postgres, hybrid search,
the API, the search step in a run, and exported code searching on its own."""

from __future__ import annotations

import io
import json
import subprocess
import sys
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from easychain.compiler import compile_flow
from easychain.knowledge import search as search_module
from easychain.knowledge.ingest import build_embeddings, ingest
from easychain.knowledge.loaders import LoadError, load_bytes, split
from easychain.knowledge.search import cite_passages, search_knowledge
from easychain.knowledge.store import KnowledgeStore
from easychain.runtime import RunOptions, run_flow
from easychain.server.app import create_app

from . import pg
from .conftest import input_step, make_spec, output_step

HELP = b"""# Help centre

## Resetting your password

Open Settings, choose Reset password and follow the link in the email. The link works for
one hour.

## Billing

Invoices are sent on the first day of each month. Refunds take five working days.

## Deleting your account

Write to support. We delete everything within thirty days.
"""


def tiny_pdf(*pages: str) -> bytes:
    """A minimal PDF with one line of text per page (no PDF library needed)."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")
    font = 3 + 2 * len(pages)
    for i, line in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 712 Td ({line}) Tj ET"
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font} 0 R >> >> >>"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n{body}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return out.getvalue()


def tiny_docx() -> bytes:
    import docx

    document = docx.Document()
    document.core_properties.title = "Returns policy"
    document.add_heading("Returns", level=1)
    document.add_paragraph("You can return anything within 30 days.")
    document.add_heading("Exchanges", level=1)
    document.add_paragraph("Exchanges are free.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Item", "Days"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Shoes", "60"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── reading and splitting ────────────────────────────────────────────────────


def test_reads_every_supported_format():
    md = load_bytes("help.md", HELP)
    assert md.kind == "markdown" and md.title == "Help centre"
    assert [s.heading for s in md.sections] == [
        "Resetting your password",
        "Billing",
        "Deleting your account",
    ]

    pdf = load_bytes("guide.pdf", tiny_pdf("First page text", "Second page text"))
    assert pdf.kind == "pdf" and [s.page for s in pdf.sections] == [1, 2]
    assert "Second page" in pdf.sections[1].text

    word = load_bytes("returns.docx", tiny_docx())
    assert word.title == "Returns policy"
    assert [s.heading for s in word.sections][:2] == ["Returns", "Exchanges"]
    assert "Shoes | 60" in word.sections[-1].text

    page = load_bytes(
        "page.html",
        b"<html><head><title>Plans</title><style>x{}</style></head><body><nav>Menu</nav>"
        b"<h1>Plans</h1><p>Starter costs ten dollars.</p><h2>Team</h2><p>Team adds sharing.</p>"
        b"<script>track()</script></body></html>",
    )
    assert page.title == "Plans"
    text = " ".join(s.text for s in page.sections)
    assert "Starter costs" in text and "Menu" not in text and "track" not in text
    assert page.sections[-1].heading == "Team"

    table = load_bytes("prices.csv", b"plan,price\nStarter,10\nTeam,25\n")
    assert table.sections[0].text == "plan: Starter; price: 10\nplan: Team; price: 25"

    with pytest.raises(LoadError, match="can't read .exe"):
        load_bytes("tool.exe", b"MZ")
    with pytest.raises(LoadError, match="no text"):
        load_bytes("scan.pdf", tiny_pdf(""))


def test_split_keeps_where_each_chunk_came_from():
    loaded = load_bytes("help.md", HELP)
    chunks = split(loaded, "help.md", chunk_size=120, chunk_overlap=20)
    assert len(chunks) >= 4
    assert all(len(c.text) <= 120 for c in chunks)
    assert [c.position for c in chunks] == list(range(len(chunks)))
    billing = next(c for c in chunks if "Invoices" in c.text)
    assert billing.heading == "Billing" and billing.source == "help.md"
    assert billing.title == "Help centre"


# ── storage and search ───────────────────────────────────────────────────────


@pytest.fixture(params=["sqlite", "postgres"])
def knowledge_url(request, tmp_path, monkeypatch):
    if request.param == "postgres":
        if not pg.available():
            pytest.skip("Postgres isn't installed")
        with pg.server() as base:
            url = pg.new_database(base)
            monkeypatch.setenv("EASYCHAIN_KNOWLEDGE_URL", url)
            yield url
            search_module.KNOWLEDGE_ENGINES.pop(url, None)
        return
    url = f"sqlite:///{tmp_path / 'kb.db'}"
    monkeypatch.setenv("EASYCHAIN_KNOWLEDGE_URL", url)
    yield url
    search_module.KNOWLEDGE_ENGINES.pop(url, None)


def _filled(store: KnowledgeStore, name: str = "Help") -> str:
    kb = store.create_base(name, "keywords", 0, chunk_size=200, chunk_overlap=20)
    for filename, data in (("help.md", HELP), ("prices.csv", b"plan,price\nStarter,10\nTeam,25\n")):
        loaded = load_bytes(filename, data)
        doc = store.add_document(kb["id"], loaded.title, filename, loaded.kind)
        ingest(store, kb["id"], doc["id"], loaded, filename)
    return kb["id"]


def test_hybrid_search_finds_and_cites_passages(knowledge_url):
    store = KnowledgeStore().setup()
    if knowledge_url.startswith("postgresql"):
        assert store.pgvector, "pgvector should be used when the extension is there"
    kb = _filled(store)
    assert store.get_base(kb)["dims"] == 256
    embeddings = build_embeddings("keywords")
    hits = search_knowledge(kb, "How do I reset my password?", embeddings, top_k=2)
    assert hits[0]["heading"] == "Resetting your password"
    assert [h["n"] for h in hits] == [1, 2]
    assert (
        search_knowledge(kb, "refunds", embeddings, top_k=1, mode="words")[0]["heading"]
        == "Billing"
    )
    assert (
        search_knowledge(kb, "Team price", embeddings, top_k=1, mode="meaning")[0]["source"]
        == "prices.csv"
    )
    assert search_knowledge(kb, "   ", embeddings) == []
    cited = cite_passages(hits[:1])
    assert cited.startswith("[1] Help centre › Resetting your password (help.md)\n")

    # Another embedding size can't be mixed into a base that already has chunks.
    with pytest.raises(ValueError, match="256-number embeddings"):
        store.ensure_dims(kb, 1536)

    docs = store.list_documents(kb)
    assert {d["status"] for d in docs} == {"ready"} and sum(d["chunks"] for d in docs) >= 4
    store.delete_document(docs[0]["id"])
    assert not search_knowledge(kb, "reset password", embeddings, mode="words")
    store.delete_base(kb)
    assert store.list_bases() == []


# ── the API ──────────────────────────────────────────────────────────────────


@pytest.fixture
def client(tmp_path, monkeypatch, fake_server):
    monkeypatch.delenv("EASYCHAIN_KNOWLEDGE_URL", raising=False)
    app = create_app(
        workspace=tmp_path / "flows",
        home=tmp_path / "home",
        static_dir=tmp_path / "web",
        database_url=f"sqlite:///{tmp_path / 'app.db'}",
        worker=False,
    )
    with TestClient(app) as test_client:
        yield test_client
    search_module.KNOWLEDGE_ENGINES.clear()


def _wait_ready(client, kb_id: str, count: int) -> list[dict]:
    for _ in range(100):
        docs = client.get(f"/api/knowledge/{kb_id}").json()["documents"]
        if len(docs) == count and all(d["status"] in ("ready", "error") for d in docs):
            return docs
        time.sleep(0.05)
    raise AssertionError(f"documents not ready: {docs}")


def test_knowledge_api(client, fake_server):
    created = client.post("/api/knowledge", json={"name": "Product docs", "chunk_size": 300})
    assert created.status_code == 201
    kb = created.json()
    assert kb["id"] == "product_docs" and kb["embedding_model"] == "keywords"
    assert kb["embedding_label"].startswith("Keywords")
    second = client.post("/api/knowledge", json={"name": "Product docs"}).json()
    assert second["id"] == "product_docs_2"

    upload = client.post(
        f"/api/knowledge/{kb['id']}/files",
        files=[
            ("files", ("help.md", HELP, "text/markdown")),
            ("files", ("tool.exe", b"MZ", "application/octet-stream")),
        ],
    )
    assert upload.status_code == 202 and len(upload.json()) == 2
    client.post(f"/api/knowledge/{kb['id']}/urls", json={"urls": [f"{fake_server.url}/pages/bees"]})
    client.post(
        f"/api/knowledge/{kb['id']}/text", json={"title": "Note", "text": "Ships in 2 days."}
    )
    docs = {d["source"]: d for d in _wait_ready(client, kb["id"], 4)}
    assert docs["help.md"]["status"] == "ready" and docs["help.md"]["chunks"] >= 3
    assert docs["tool.exe"]["status"] == "error" and "can't read" in docs["tool.exe"]["error"]
    assert docs[f"{fake_server.url}/pages/bees"]["title"] == "Honey bees"

    hits = client.post(
        f"/api/knowledge/{kb['id']}/search", json={"query": "what do worker bees collect"}
    ).json()["hits"]
    assert "nectar" in hits[0]["text"]

    chunks = client.get(
        f"/api/knowledge/{kb['id']}/documents/{docs['help.md']['id']}/chunks"
    ).json()
    assert chunks[0]["position"] == 0 and chunks[0]["title"] == "Help centre"

    listed = client.get("/api/knowledge").json()
    # The help centre the Support bot template uses is there from the start.
    assert {b["id"] for b in listed} == {"product_docs", "product_docs_2", "help_centre"}
    assert next(b for b in listed if b["id"] == "product_docs")["documents"] == 4

    preview = client.post(
        "/api/knowledge-preview",
        files={"file": ("help.md", HELP, "text/markdown")},
        data={"chunk_size": "150", "chunk_overlap": "0"},
    ).json()
    assert preview["title"] == "Help centre" and preview["count"] >= 3
    assert all(len(c["text"]) <= 150 for c in preview["chunks"])
    bad = client.post("/api/knowledge-preview", data={"url": "not a url"})
    assert bad.status_code == 422 and "web address" in bad.json()["message"]

    client.delete(f"/api/knowledge/{kb['id']}/documents/{docs['help.md']['id']}")
    assert len(client.get(f"/api/knowledge/{kb['id']}").json()["documents"]) == 3
    assert (
        client.patch(f"/api/knowledge/{kb['id']}", json={"name": "Docs"}).json()["name"] == "Docs"
    )
    assert client.delete(f"/api/knowledge/{second['id']}").status_code == 200
    assert client.get(f"/api/knowledge/{second['id']}").status_code == 404


def test_checks_warn_about_a_missing_base_or_another_model(client):
    kb = client.post("/api/knowledge", json={"name": "Docs"}).json()
    spec = make_spec(
        [
            input_step("question"),
            {"id": "search", "type": "knowledge_search", "settings": {"knowledge_base": "nope"}},
            {
                "id": "search2",
                "type": "knowledge_search",
                "settings": {
                    "knowledge_base": kb["id"],
                    "embedding_model": "openai:text-embedding-3-small",
                    "save_as": "c2",
                    "sources_as": "s2",
                },
            },
            output_step("context"),
        ],
        [("input", "search"), ("search", "search2"), ("search2", "output")],
    )
    issues = client.post("/api/check", json=json.loads(spec.model_dump_json(by_alias=True))).json()[
        "issues"
    ]
    codes = {(i["code"], i["step"]) for i in issues}
    assert ("knowledge_base_missing", "search") in codes
    mismatch = next(i for i in issues if i["code"] == "knowledge_model_mismatch")
    assert mismatch["fix"]["params"] == {"key": "embedding_model", "value": "keywords"}


# ── in a run, and in exported code ───────────────────────────────────────────


def _answer_flow(kb_id: str):
    return make_spec(
        [
            input_step("question"),
            {
                "id": "search",
                "type": "knowledge_search",
                "settings": {"knowledge_base": kb_id, "top_k": 2},
            },
            {
                "id": "write_prompt",
                "type": "instructions",
                "settings": {
                    "system": "Answer from the passages and cite them like [1].",
                    "user": "Question: {question}\n\nPassages:\n\n{context}",
                },
            },
            {"id": "answer_it", "type": "ai_model", "settings": {"save_as": "answer"}},
            output_step("answer", "sources"),
        ],
        [
            ("input", "search"),
            ("search", "write_prompt"),
            ("write_prompt", "answer_it"),
            ("answer_it", "output"),
        ],
    )


async def test_search_step_in_a_run_gives_passages_and_sources(knowledge_url):
    store = KnowledgeStore().setup()
    kb = _filled(store)
    final, events = await run_flow(
        _answer_flow(kb), {"question": "How long do refunds take?"}, RunOptions(stand_in=True)
    )
    assert final["status"] == "ok", final
    sources = final["output"]["sources"]
    assert sources[0]["heading"] == "Billing" and sources[0]["n"] == 1
    searched = next(e for e in events if e["type"] == "step_finished" and e["step"] == "search")
    assert searched["output"]["context"].startswith("[1] Help centre › Billing")


def test_exported_code_searches_on_its_own(tmp_path, monkeypatch):
    url = f"sqlite:///{tmp_path / 'kb.db'}"
    monkeypatch.setenv("EASYCHAIN_KNOWLEDGE_URL", url)
    kb = _filled(KnowledgeStore().setup())
    spec = make_spec(
        [
            input_step("question"),
            {"id": "search", "type": "knowledge_search", "settings": {"knowledge_base": kb}},
            output_step("context", "sources"),
        ],
        [("input", "search"), ("search", "output")],
    )
    compiled = compile_flow(spec)
    assert "sqlalchemy>=2.0.36" in compiled.requirements and "numpy>=2" in compiled.requirements
    module = tmp_path / f"{compiled.module_name}.py"
    module.write_text(compiled.source)
    env = {"EASYCHAIN_KNOWLEDGE_URL": url, "PATH": "/usr/bin:/bin"}
    result = subprocess.run(
        [sys.executable, str(module), json.dumps({"question": "delete my account"})],
        capture_output=True,
        text=True,
        env=env,
        cwd=Path(module).parent,
        timeout=120,
    )
    assert result.returncode == 0, result.stderr
    out = json.loads(result.stdout)
    assert out["sources"][0]["heading"] == "Deleting your account"
    search_module.KNOWLEDGE_ENGINES.pop(url, None)
