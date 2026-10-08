"""Reading files and web pages into sections, and splitting sections into chunks.

Supported: PDF (pypdf), Word (.docx), HTML, Markdown, CSV and plain text, from an
upload or a URL. Each section keeps where it came from (page or heading), so a
chunk can be cited precisely.
"""

from __future__ import annotations

import codecs
import csv
import io
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import PurePosixPath

from langchain_text_splitters import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter

from ..fetch import FetchError, TooLarge, fetch, size_text
from .store import ChunkDraft

KINDS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".html": "html",
    ".htm": "html",
    ".md": "markdown",
    ".markdown": "markdown",
    ".csv": "csv",
    ".txt": "text",
    ".text": "text",
    ".json": "text",
    ".yaml": "text",
    ".yml": "text",
}
MAX_BYTES = 50 * 1024 * 1024


class LoadError(ValueError):
    """A file or page that can't be read; the message is shown to people."""


@dataclass
class Section:
    text: str
    page: int | None = None
    heading: str | None = None


@dataclass
class Loaded:
    kind: str
    title: str
    sections: list[Section]


def kind_for(name: str, content_type: str | None = None) -> str:
    suffix = PurePosixPath(name.lower().split("?")[0]).suffix
    if suffix in KINDS:
        return KINDS[suffix]
    ctype = (content_type or "").split(";")[0].strip().lower()
    return {
        "application/pdf": "pdf",
        "text/html": "html",
        "text/markdown": "markdown",
        "text/csv": "csv",
        "text/plain": "text",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    }.get(ctype, "")


def _utf16_without_bom(data: bytes) -> str | None:
    """The UTF-16 codec for text saved without a byte order mark, or None.

    Text in UTF-16 has a zero byte beside nearly every Latin letter; other text has almost
    none. (Simply trying UTF-16 would "succeed" on most files, giving garbage.)
    """
    sample = data[:4096]
    pairs = len(sample) // 2
    if pairs < 2 or len(data) % 2:
        return None
    if sample[1::2].count(0) > pairs * 0.3 and sample[0::2].count(0) < pairs * 0.05:
        return "utf-16-le"
    if sample[0::2].count(0) > pairs * 0.3 and sample[1::2].count(0) < pairs * 0.05:
        return "utf-16-be"
    return None


def _decode(data: bytes) -> str:
    """Text from bytes: UTF-8, UTF-16 or UTF-32 (with a byte order mark, or UTF-16 that
    clearly is), else Windows-1252 (Western European), which also covers Latin-1."""
    for bom, encoding in (
        (codecs.BOM_UTF32_LE, "utf-32"),
        (codecs.BOM_UTF32_BE, "utf-32"),
        (codecs.BOM_UTF8, "utf-8-sig"),
        (codecs.BOM_UTF16_LE, "utf-16"),
        (codecs.BOM_UTF16_BE, "utf-16"),
    ):
        if data.startswith(bom):
            return data.decode(encoding, errors="replace")
    utf16 = _utf16_without_bom(data)
    if utf16:
        return data.decode(utf16, errors="replace")
    for encoding in ("utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")  # every byte is a Latin-1 character


def _stem(name: str) -> str:
    return (
        PurePosixPath(name.split("?")[0]).stem.replace("-", " ").replace("_", " ").strip() or name
    )


# ── formats ──────────────────────────────────────────────────────────────────


def _pdf(data: bytes, name: str) -> Loaded:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [
            Section(page.extract_text() or "", page=n) for n, page in enumerate(reader.pages, 1)
        ]
    except (PdfReadError, ValueError) as exc:
        raise LoadError(f"{name} isn't a PDF I can read ({exc}).") from exc
    title = ""
    try:
        title = (reader.metadata.title or "") if reader.metadata else ""
    except Exception:  # damaged metadata is not worth failing for
        title = ""
    if not any(p.text.strip() for p in pages):
        raise LoadError(
            f"{name} has no text to read (it may be scanned images). Text recognition (OCR) "
            "isn't supported yet."
        )
    return Loaded("pdf", title.strip() or _stem(name), [p for p in pages if p.text.strip()])


def _docx(data: bytes, name: str) -> Loaded:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise LoadError(f"{name} isn't a Word document I can read.") from exc
    sections: list[Section] = []
    heading: str | None = None
    lines: list[str] = []

    def flush() -> None:
        if lines:
            sections.append(Section("\n".join(lines), heading=heading))
            lines.clear()

    for para in document.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        if style.lower().startswith("heading") and para.text.strip():
            flush()
            heading = para.text.strip()
            lines.append(heading)
        elif para.text.strip():
            lines.append(para.text.strip())
    flush()
    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            sections.append(Section("\n".join(rows), heading=heading))
    title = (document.core_properties.title or "").strip()
    return Loaded("docx", title or _stem(name), sections)


class _Readable(HTMLParser):
    """The words people read on a page, split at headings (no menus, scripts or styles)."""

    SKIP = {"script", "style", "nav", "header", "footer", "noscript", "svg", "form", "aside"}
    BLOCKS = {"p", "div", "li", "br", "tr", "section", "article", "pre", "blockquote", "td", "th"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.skip = 0
        self.title = ""
        self.in_title = False
        self.heading_tag: str | None = None
        self.heading_text: list[str] = []
        self.sections: list[Section] = []
        self.current: list[str] = []
        self.heading: str | None = None

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag in self.SKIP:
            self.skip += 1
        elif tag == "title":
            self.in_title = True
        elif re.fullmatch(r"h[1-6]", tag) and not self.skip:
            self._flush()
            self.heading_tag, self.heading_text = tag, []
        elif tag in self.BLOCKS:
            self.current.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        elif tag == "title":
            self.in_title = False
        elif tag == self.heading_tag:
            self.heading = " ".join("".join(self.heading_text).split())
            self.current.append(self.heading + "\n")
            self.heading_tag = None
        elif tag in self.BLOCKS:
            self.current.append("\n")

    def handle_data(self, data: str) -> None:
        if self.in_title:
            self.title += data
        elif self.skip:
            return
        elif self.heading_tag:
            self.heading_text.append(data)
        else:
            self.current.append(data)

    def _flush(self) -> None:
        body = re.sub(r"[ \t]+", " ", "".join(self.current))
        body = re.sub(r"\n\s*\n+", "\n\n", body).strip()
        if body:
            self.sections.append(Section(body, heading=self.heading))
        self.current = []

    def done(self) -> list[Section]:
        self._flush()
        return self.sections


def _html(data: bytes, name: str) -> Loaded:
    parser = _Readable()
    parser.feed(_decode(data))
    sections = parser.done()
    if not sections:
        raise LoadError(f"{name} has no readable text.")
    return Loaded("html", " ".join(parser.title.split()) or _stem(name), sections)


def _markdown(data: bytes, name: str) -> Loaded:
    content = _decode(data)
    splitter = MarkdownHeaderTextSplitter(
        headers_to_split_on=[("#", "h1"), ("##", "h2"), ("###", "h3")], strip_headers=False
    )
    sections = []
    for doc in splitter.split_text(content):
        heading = doc.metadata.get("h3") or doc.metadata.get("h2") or doc.metadata.get("h1")
        if doc.page_content.strip():
            sections.append(Section(doc.page_content, heading=heading))
    first = re.search(r"^#\s+(.+)$", content, flags=re.M)
    return Loaded("markdown", first.group(1).strip() if first else _stem(name), sections)


def _csv(data: bytes, name: str) -> Loaded:
    reader = csv.reader(io.StringIO(_decode(data)))
    try:
        rows = [row for row in reader if any(cell.strip() for cell in row)]
    except csv.Error as exc:  # e.g. a value over 131,072 characters
        raise LoadError(
            f"{name} isn't a CSV file I can read ({exc}). Check for a quote (\") that's never "
            "closed, or save the file as text (.txt) instead."
        ) from exc
    if not rows:
        raise LoadError(f"{name} is empty.")
    header, body = rows[0], rows[1:]
    lines = [
        "; ".join(
            f"{h.strip()}: {v.strip()}" for h, v in zip(header, row, strict=False) if v.strip()
        )
        for row in body
    ]
    return Loaded("csv", _stem(name), [Section("\n".join(lines))])


def _text(data: bytes, name: str) -> Loaded:
    content = _decode(data)
    if not content.strip():
        raise LoadError(f"{name} is empty.")
    return Loaded("text", _stem(name), [Section(content)])


READERS = {
    "pdf": _pdf,
    "docx": _docx,
    "html": _html,
    "markdown": _markdown,
    "csv": _csv,
    "text": _text,
}


def _read(kind: str, data: bytes, name: str) -> Loaded:
    """Read with the reader for ``kind``; whatever goes wrong becomes a LoadError."""
    try:
        return READERS[kind](data, name)
    except LoadError:
        raise
    except Exception as exc:  # a damaged or unusual file: say so, don't crash
        what = {
            "pdf": "a PDF",
            "docx": "a Word document",
            "html": "a web page",
            "markdown": "a Markdown file",
            "csv": "a CSV file",
        }.get(kind, "text")
        reason = " ".join(str(exc).split())[:300] or type(exc).__name__
        raise LoadError(
            f"I couldn't read {name} as {what} ({reason}). Check that the file isn't "
            "damaged, or save it in another format and add it again."
        ) from exc


def load_bytes(name: str, data: bytes, content_type: str | None = None) -> Loaded:
    if len(data) > MAX_BYTES:
        raise LoadError(
            f"{name} is larger than {size_text(MAX_BYTES)}. Split it into smaller files."
        )
    kind = kind_for(name, content_type)
    if not kind:
        raise LoadError(
            f"I can't read {PurePosixPath(name).suffix or 'this kind of'} files yet. "
            "Use PDF, Word (.docx), HTML, Markdown, CSV or text."
        )
    return _read(kind, data, name)


def load_url(url: str, timeout: float = 30) -> Loaded:
    if not re.match(r"^https?://", url):
        raise LoadError(f"“{url}” isn't a web address (it should start with https://).")
    try:
        page = fetch(
            url,
            limit=MAX_BYTES,
            timeout=timeout,
            headers={"User-Agent": "Mozilla/5.0 (compatible; EasyChain/0.3)"},
        )
    except TooLarge as exc:
        raise LoadError(f"{exc} Download it and add the parts you need as files.") from None
    except FetchError as exc:
        raise LoadError(str(exc)) from None
    kind = kind_for(url, page.content_type) or "html"
    return _read(kind, page.content, url)


# ── chunks ───────────────────────────────────────────────────────────────────


def split(
    loaded: Loaded, source: str, chunk_size: int = 1000, chunk_overlap: int = 150
) -> list[ChunkDraft]:
    """Split sections into chunks of about ``chunk_size`` characters, overlapping a little."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size, chunk_overlap=min(chunk_overlap, chunk_size // 2)
    )
    chunks: list[ChunkDraft] = []
    for section in loaded.sections:
        for raw in splitter.split_text(section.text):
            piece = "\n".join(line.rstrip() for line in raw.splitlines())
            if piece.strip():
                chunks.append(
                    ChunkDraft(
                        text=piece.strip(),
                        position=len(chunks),
                        title=loaded.title,
                        source=source,
                        page=section.page,
                        heading=section.heading,
                    )
                )
    return chunks
